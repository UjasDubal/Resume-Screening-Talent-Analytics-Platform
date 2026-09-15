"""
Synthetic Resume Generator
===========================
Generates 200 synthetic resumes as a mix of PDF and DOCX files across 5 job families:
  - Software Engineering
  - Data Analysis
  - Marketing
  - Sales
  - Design

Uses Faker for PII, hardcoded skill/role templates for realism.
PDF layouts: single-column, two-column, table-based (via reportlab).
DOCX layouts: standard document format (via python-docx).

No Groq API dependency — fully offline generation.
"""

import os
import sys
import json
import random
import math
from datetime import datetime, timedelta
from pathlib import Path

from faker import Faker
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    Frame, PageTemplate, BaseDocTemplate
)
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


fake = Faker()
Faker.seed(42)
random.seed(42)

# ---------------------------------------------------------------------------
# Job Family Definitions
# ---------------------------------------------------------------------------

JOB_FAMILIES = {
    "Software Engineering": {
        "titles": [
            "Software Engineer", "Senior Software Engineer", "Backend Developer",
            "Frontend Developer", "Full Stack Developer", "DevOps Engineer",
            "Site Reliability Engineer", "Software Architect", "Mobile Developer",
            "Cloud Engineer", "Platform Engineer"
        ],
        "skills": [
            "Python", "Java", "JavaScript", "TypeScript", "C++", "Go", "Rust",
            "React", "Angular", "Vue.js", "Node.js", "Django", "Flask", "FastAPI",
            "Spring Boot", "Docker", "Kubernetes", "AWS", "Azure", "GCP",
            "PostgreSQL", "MySQL", "MongoDB", "Redis", "Git", "CI/CD",
            "REST APIs", "GraphQL", "Microservices", "Linux", "Terraform",
            "Jenkins", "GitHub Actions"
        ],
        "education": [
            "B.S. Computer Science", "M.S. Computer Science",
            "B.S. Software Engineering", "B.S. Information Technology",
            "M.S. Software Engineering", "B.E. Computer Engineering"
        ],
        "certifications": [
            "AWS Certified Solutions Architect", "Google Cloud Professional",
            "Kubernetes Administrator (CKA)", "Azure Developer Associate",
            "HashiCorp Terraform Associate"
        ]
    },
    "Data Analysis": {
        "titles": [
            "Data Analyst", "Senior Data Analyst", "Business Intelligence Analyst",
            "Data Scientist", "Machine Learning Engineer", "Analytics Engineer",
            "Quantitative Analyst", "Research Analyst", "Data Engineer",
            "Statistical Analyst"
        ],
        "skills": [
            "Python", "R", "SQL", "Tableau", "Power BI", "Excel",
            "Pandas", "NumPy", "Scikit-learn", "TensorFlow", "PyTorch",
            "Spark", "Hadoop", "Airflow", "dbt", "Snowflake", "BigQuery",
            "Statistics", "A/B Testing", "Data Visualization", "ETL",
            "Machine Learning", "Deep Learning", "NLP", "Feature Engineering",
            "Jupyter", "Git", "AWS", "Docker"
        ],
        "education": [
            "B.S. Statistics", "M.S. Data Science", "B.S. Mathematics",
            "M.S. Statistics", "B.S. Computer Science", "Ph.D. Statistics",
            "M.S. Analytics", "B.S. Economics"
        ],
        "certifications": [
            "Google Data Analytics Certificate", "AWS Data Analytics Specialty",
            "Tableau Desktop Specialist", "Microsoft Power BI Data Analyst",
            "IBM Data Science Professional"
        ]
    },
    "Marketing": {
        "titles": [
            "Marketing Manager", "Digital Marketing Specialist",
            "Content Marketing Manager", "SEO Specialist", "Social Media Manager",
            "Growth Marketing Manager", "Brand Manager", "Marketing Analyst",
            "Email Marketing Specialist", "Product Marketing Manager"
        ],
        "skills": [
            "SEO", "SEM", "Google Analytics", "Google Ads", "Facebook Ads",
            "Content Strategy", "Copywriting", "Social Media Marketing",
            "Email Marketing", "Marketing Automation", "HubSpot", "Salesforce",
            "A/B Testing", "Brand Strategy", "Market Research", "CRM",
            "Adobe Creative Suite", "Canva", "WordPress", "Mailchimp",
            "Campaign Management", "Lead Generation", "Conversion Optimization",
            "Public Relations", "Event Marketing"
        ],
        "education": [
            "B.A. Marketing", "M.B.A. Marketing", "B.A. Communications",
            "B.S. Business Administration", "M.A. Digital Marketing",
            "B.A. English", "B.A. Journalism"
        ],
        "certifications": [
            "Google Ads Certification", "HubSpot Inbound Marketing",
            "Facebook Blueprint", "Google Analytics Individual Qualification",
            "Hootsuite Social Marketing"
        ]
    },
    "Sales": {
        "titles": [
            "Sales Representative", "Account Executive", "Sales Manager",
            "Business Development Representative", "Enterprise Sales Executive",
            "Regional Sales Manager", "Inside Sales Representative",
            "Sales Operations Analyst", "Key Account Manager",
            "VP of Sales"
        ],
        "skills": [
            "Salesforce", "CRM Management", "Cold Calling", "Lead Generation",
            "Pipeline Management", "Negotiation", "Contract Negotiation",
            "Account Management", "B2B Sales", "B2C Sales",
            "Territory Management", "Revenue Forecasting", "Presentation Skills",
            "Relationship Building", "Solution Selling", "Consultative Selling",
            "Prospecting", "Closing Deals", "Cross-Selling", "Upselling",
            "HubSpot", "Sales Analytics", "Quota Attainment",
            "Customer Retention", "Strategic Planning"
        ],
        "education": [
            "B.A. Business Administration", "M.B.A.",
            "B.S. Marketing", "B.A. Communications",
            "B.S. Business Management", "B.A. Economics"
        ],
        "certifications": [
            "Salesforce Administrator", "HubSpot Sales Software",
            "Certified Professional Sales Leader (CPSL)",
            "Strategic Selling Certification", "SPIN Selling Certified"
        ]
    },
    "Design": {
        "titles": [
            "UI/UX Designer", "Senior Product Designer", "Visual Designer",
            "Graphic Designer", "Interaction Designer", "UX Researcher",
            "Design Lead", "Creative Director", "Web Designer",
            "Motion Designer"
        ],
        "skills": [
            "Figma", "Sketch", "Adobe XD", "Adobe Photoshop",
            "Adobe Illustrator", "InVision", "Zeplin", "Wireframing",
            "Prototyping", "User Research", "Usability Testing",
            "Design Systems", "Typography", "Color Theory",
            "Responsive Design", "HTML", "CSS", "JavaScript",
            "After Effects", "Blender", "3D Modeling",
            "Information Architecture", "Accessibility Design",
            "Design Thinking", "Brand Identity"
        ],
        "education": [
            "B.F.A. Graphic Design", "B.A. Visual Arts",
            "M.F.A. Design", "B.S. Human-Computer Interaction",
            "B.A. Industrial Design", "B.A. Fine Arts"
        ],
        "certifications": [
            "Google UX Design Certificate", "Interaction Design Foundation",
            "Adobe Certified Expert", "Nielsen Norman Group UX Certification",
            "Certified Usability Analyst"
        ]
    }
}

# Common action verbs for bullet points
ACTION_VERBS = [
    "Developed", "Implemented", "Designed", "Led", "Managed", "Optimized",
    "Created", "Built", "Delivered", "Improved", "Increased", "Reduced",
    "Spearheaded", "Collaborated", "Launched", "Automated", "Streamlined",
    "Mentored", "Analyzed", "Architected", "Established", "Pioneered"
]

INDUSTRIES = [
    "Technology", "Finance", "Healthcare", "E-commerce", "Education",
    "Media", "Consulting", "Manufacturing", "Retail", "Telecommunications",
    "Automotive", "Energy", "Real Estate", "Insurance", "Logistics"
]


def generate_candidate_data(candidate_id: int) -> dict:
    """Generate a single synthetic candidate profile."""
    family_name = random.choice(list(JOB_FAMILIES.keys()))
    family = JOB_FAMILIES[family_name]

    name = fake.name()
    email = fake.email()
    phone = fake.phone_number()
    city = fake.city()
    state = fake.state_abbr()
    linkedin = f"linkedin.com/in/{name.lower().replace(' ', '-')}-{random.randint(100,999)}"

    # Experience: 0-20 years
    experience_years = random.choices(
        range(0, 21),
        weights=[3, 5, 7, 8, 9, 10, 10, 9, 8, 7, 6, 5, 4, 3, 2, 2, 1, 1, 1, 1, 1],
        k=1
    )[0]

    # Number of skills: 5-12
    num_skills = random.randint(5, min(12, len(family["skills"])))
    skills = random.sample(family["skills"], num_skills)

    # Sometimes add cross-family skills
    if random.random() > 0.6:
        other_family = random.choice([f for f in JOB_FAMILIES.keys() if f != family_name])
        cross_skills = random.sample(JOB_FAMILIES[other_family]["skills"], random.randint(1, 3))
        skills.extend(cross_skills)

    education = random.choice(family["education"])
    university = fake.company() + " University" if random.random() > 0.3 else random.choice([
        "MIT", "Stanford University", "UC Berkeley", "Carnegie Mellon University",
        "Georgia Tech", "University of Michigan", "University of Texas at Austin",
        "University of Washington", "Purdue University", "Penn State University"
    ])

    grad_year = datetime.now().year - experience_years - random.randint(0, 4)

    # Work experience entries
    num_jobs = min(experience_years // 2 + 1, random.randint(1, 5))
    work_experience = []
    current_year = datetime.now().year

    for j in range(num_jobs):
        title = random.choice(family["titles"])
        company = fake.company()
        industry = random.choice(INDUSTRIES)

        if j == 0:
            end_date = "Present"
            end_year = current_year
        else:
            end_year = current_year - sum(
                we.get("duration_years", 1) for we in work_experience
            )
            end_date = f"{random.choice(['Jan','Mar','Jun','Sep'])} {end_year}"

        duration = random.randint(1, max(1, experience_years // num_jobs + 1))
        start_year = end_year - duration
        start_date = f"{random.choice(['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'])} {start_year}"

        # Generate 2-4 bullet points
        bullets = []
        for _ in range(random.randint(2, 4)):
            verb = random.choice(ACTION_VERBS)
            skill_mention = random.choice(skills) if random.random() > 0.3 else ""
            metric = ""
            if random.random() > 0.4:
                metric = random.choice([
                    f" resulting in {random.randint(10,60)}% improvement",
                    f" serving {random.randint(1,50)}K+ users",
                    f" reducing costs by ${random.randint(10,500)}K annually",
                    f" across {random.randint(3,15)} teams",
                    f" increasing revenue by {random.randint(5,40)}%",
                    f" processing {random.randint(1,100)}M+ records daily",
                ])

            bullet = f"{verb} {fake.bs()}"
            if skill_mention:
                bullet += f" using {skill_mention}"
            bullet += metric
            bullets.append(bullet)

        work_experience.append({
            "title": title,
            "company": company,
            "industry": industry,
            "start_date": start_date,
            "end_date": end_date,
            "duration_years": duration,
            "bullets": bullets
        })

    # Certifications (0-2)
    num_certs = random.choices([0, 1, 2], weights=[40, 40, 20], k=1)[0]
    certifications = random.sample(
        family["certifications"], min(num_certs, len(family["certifications"]))
    )

    # Summary statement
    summary = (
        f"Results-driven {random.choice(family['titles']).lower()} with "
        f"{experience_years}+ years of experience in {family_name.lower()}. "
        f"Proficient in {', '.join(skills[:3])} and {skills[3] if len(skills) > 3 else 'more'}. "
        f"Passionate about delivering high-quality solutions and driving business impact."
    )

    return {
        "candidate_id": candidate_id,
        "name": name,
        "email": email,
        "phone": phone,
        "city": city,
        "state": state,
        "linkedin": linkedin,
        "summary": summary,
        "job_family": family_name,
        "experience_years": experience_years,
        "skills": skills,
        "education": education,
        "university": university,
        "grad_year": grad_year,
        "work_experience": work_experience,
        "certifications": certifications,
    }


# ---------------------------------------------------------------------------
# PDF Generators (3 layout variants)
# ---------------------------------------------------------------------------

def _create_pdf_single_column(candidate: dict, filepath: str):
    """Standard single-column resume PDF."""
    doc = SimpleDocTemplate(
        filepath, pagesize=letter,
        topMargin=0.5 * inch, bottomMargin=0.5 * inch,
        leftMargin=0.75 * inch, rightMargin=0.75 * inch
    )
    styles = getSampleStyleSheet()

    # Custom styles
    name_style = ParagraphStyle(
        "Name", parent=styles["Title"], fontSize=18,
        spaceAfter=2, textColor=colors.HexColor("#1a1a2e")
    )
    contact_style = ParagraphStyle(
        "Contact", parent=styles["Normal"], fontSize=9,
        textColor=colors.grey, alignment=TA_CENTER, spaceAfter=10
    )
    section_style = ParagraphStyle(
        "Section", parent=styles["Heading2"], fontSize=12,
        textColor=colors.HexColor("#16213e"), spaceAfter=4,
        spaceBefore=10, borderWidth=0, borderPadding=0
    )
    body_style = ParagraphStyle(
        "Body", parent=styles["Normal"], fontSize=10,
        spaceAfter=2, leading=13
    )
    bullet_style = ParagraphStyle(
        "Bullet", parent=styles["Normal"], fontSize=9,
        leftIndent=20, spaceAfter=1, leading=12,
        bulletIndent=10
    )

    elements = []

    # Header
    elements.append(Paragraph(candidate["name"], name_style))
    contact_line = (
        f"{candidate['email']} | {candidate['phone']} | "
        f"{candidate['city']}, {candidate['state']} | {candidate['linkedin']}"
    )
    elements.append(Paragraph(contact_line, contact_style))
    elements.append(Spacer(1, 6))

    # Summary
    elements.append(Paragraph("PROFESSIONAL SUMMARY", section_style))
    elements.append(Paragraph(candidate["summary"], body_style))

    # Skills
    elements.append(Paragraph("SKILLS", section_style))
    skills_text = " • ".join(candidate["skills"])
    elements.append(Paragraph(skills_text, body_style))

    # Experience
    elements.append(Paragraph("WORK EXPERIENCE", section_style))
    for job in candidate["work_experience"]:
        job_header = (
            f"<b>{job['title']}</b> — {job['company']} "
            f"({job['start_date']} – {job['end_date']})"
        )
        elements.append(Paragraph(job_header, body_style))
        for bullet in job["bullets"]:
            elements.append(Paragraph(f"• {bullet}", bullet_style))
        elements.append(Spacer(1, 4))

    # Education
    elements.append(Paragraph("EDUCATION", section_style))
    edu_text = f"<b>{candidate['education']}</b> — {candidate['university']} ({candidate['grad_year']})"
    elements.append(Paragraph(edu_text, body_style))

    # Certifications
    if candidate["certifications"]:
        elements.append(Paragraph("CERTIFICATIONS", section_style))
        for cert in candidate["certifications"]:
            elements.append(Paragraph(f"• {cert}", bullet_style))

    doc.build(elements)


def _create_pdf_two_column(candidate: dict, filepath: str):
    """Two-column layout: sidebar (skills/contact) + main content."""
    from reportlab.platypus import FrameBreak

    page_w, page_h = letter
    sidebar_w = 2.2 * inch
    main_w = page_w - sidebar_w - 1.0 * inch  # margins

    # Create a doc with two-column page template
    doc = BaseDocTemplate(
        filepath, pagesize=letter,
        topMargin=0.5 * inch, bottomMargin=0.5 * inch,
        leftMargin=0.4 * inch, rightMargin=0.4 * inch
    )

    sidebar_frame = Frame(
        0.4 * inch, 0.5 * inch, sidebar_w, page_h - 1.0 * inch,
        id="sidebar", showBoundary=0
    )
    main_frame = Frame(
        0.4 * inch + sidebar_w + 0.2 * inch, 0.5 * inch,
        main_w, page_h - 1.0 * inch,
        id="main", showBoundary=0
    )

    doc.addPageTemplates([
        PageTemplate(id="twocol", frames=[sidebar_frame, main_frame])
    ])

    styles = getSampleStyleSheet()
    sidebar_head = ParagraphStyle(
        "SideHead", parent=styles["Heading3"], fontSize=10,
        textColor=colors.HexColor("#0f3460"), spaceAfter=4, spaceBefore=8
    )
    sidebar_body = ParagraphStyle(
        "SideBody", parent=styles["Normal"], fontSize=8,
        spaceAfter=2, leading=11
    )
    name_style = ParagraphStyle(
        "Name", parent=styles["Title"], fontSize=16,
        spaceAfter=4, textColor=colors.HexColor("#1a1a2e")
    )
    section_style = ParagraphStyle(
        "Section", parent=styles["Heading2"], fontSize=11,
        textColor=colors.HexColor("#16213e"), spaceAfter=4, spaceBefore=8
    )
    body_style = ParagraphStyle(
        "Body", parent=styles["Normal"], fontSize=9,
        spaceAfter=2, leading=12
    )
    bullet_style = ParagraphStyle(
        "Bullet", parent=styles["Normal"], fontSize=8,
        leftIndent=15, spaceAfter=1, leading=11
    )

    elements = []

    # --- Sidebar content ---
    elements.append(Paragraph("CONTACT", sidebar_head))
    elements.append(Paragraph(candidate["email"], sidebar_body))
    elements.append(Paragraph(candidate["phone"], sidebar_body))
    elements.append(Paragraph(f"{candidate['city']}, {candidate['state']}", sidebar_body))
    elements.append(Paragraph(candidate["linkedin"], sidebar_body))

    elements.append(Paragraph("SKILLS", sidebar_head))
    for skill in candidate["skills"]:
        elements.append(Paragraph(f"• {skill}", sidebar_body))

    if candidate["certifications"]:
        elements.append(Paragraph("CERTIFICATIONS", sidebar_head))
        for cert in candidate["certifications"]:
            elements.append(Paragraph(f"• {cert}", sidebar_body))

    elements.append(Paragraph("EDUCATION", sidebar_head))
    elements.append(Paragraph(candidate["education"], sidebar_body))
    elements.append(Paragraph(candidate["university"], sidebar_body))
    elements.append(Paragraph(str(candidate["grad_year"]), sidebar_body))

    # Switch to main frame
    elements.append(FrameBreak())

    # --- Main content ---
    elements.append(Paragraph(candidate["name"], name_style))
    elements.append(Spacer(1, 4))

    elements.append(Paragraph("SUMMARY", section_style))
    elements.append(Paragraph(candidate["summary"], body_style))

    elements.append(Paragraph("EXPERIENCE", section_style))
    for job in candidate["work_experience"]:
        job_header = (
            f"<b>{job['title']}</b> — {job['company']} "
            f"({job['start_date']} – {job['end_date']})"
        )
        elements.append(Paragraph(job_header, body_style))
        for bullet in job["bullets"]:
            elements.append(Paragraph(f"• {bullet}", bullet_style))
        elements.append(Spacer(1, 3))

    doc.build(elements)


def _create_pdf_table_layout(candidate: dict, filepath: str):
    """Table-based resume layout — uses Tables for structured sections."""
    doc = SimpleDocTemplate(
        filepath, pagesize=letter,
        topMargin=0.5 * inch, bottomMargin=0.5 * inch,
        leftMargin=0.6 * inch, rightMargin=0.6 * inch
    )
    styles = getSampleStyleSheet()

    name_style = ParagraphStyle(
        "Name", parent=styles["Title"], fontSize=18,
        textColor=colors.HexColor("#1a1a2e"), spaceAfter=2
    )
    section_style = ParagraphStyle(
        "Section", parent=styles["Heading2"], fontSize=11,
        textColor=colors.white, spaceAfter=0, spaceBefore=0
    )
    body_style = ParagraphStyle(
        "Body", parent=styles["Normal"], fontSize=9,
        spaceAfter=2, leading=12
    )
    small_style = ParagraphStyle(
        "Small", parent=styles["Normal"], fontSize=8,
        spaceAfter=1, leading=10
    )

    elements = []
    page_w = letter[0] - 1.2 * inch

    # Header table: Name | Contact info
    header_data = [
        [
            Paragraph(candidate["name"], name_style),
            Paragraph(
                f"{candidate['email']}<br/>{candidate['phone']}<br/>"
                f"{candidate['city']}, {candidate['state']}",
                ParagraphStyle("ContactR", parent=styles["Normal"],
                               fontSize=9, alignment=TA_RIGHT)
            )
        ]
    ]
    header_table = Table(header_data, colWidths=[page_w * 0.6, page_w * 0.4])
    header_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBELOW", (0, 0), (-1, 0), 1.5, colors.HexColor("#16213e")),
    ]))
    elements.append(header_table)
    elements.append(Spacer(1, 8))

    # Summary
    elements.append(Paragraph(candidate["summary"], body_style))
    elements.append(Spacer(1, 6))

    # Skills table (3 columns)
    section_header = Table(
        [[Paragraph("TECHNICAL SKILLS", section_style)]],
        colWidths=[page_w]
    )
    section_header.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#16213e")),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]))
    elements.append(section_header)

    skills = candidate["skills"]
    cols = 3
    rows_needed = math.ceil(len(skills) / cols)
    skill_grid = []
    for r in range(rows_needed):
        row = []
        for c in range(cols):
            idx = r * cols + c
            if idx < len(skills):
                row.append(Paragraph(f"• {skills[idx]}", small_style))
            else:
                row.append("")
        skill_grid.append(row)

    skill_table = Table(skill_grid, colWidths=[page_w / 3] * 3)
    skill_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e0e0e0")),
    ]))
    elements.append(skill_table)
    elements.append(Spacer(1, 6))

    # Experience table
    exp_header = Table(
        [[Paragraph("WORK EXPERIENCE", section_style)]],
        colWidths=[page_w]
    )
    exp_header.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#16213e")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]))
    elements.append(exp_header)

    for job in candidate["work_experience"]:
        exp_data = [
            [
                Paragraph(f"<b>{job['title']}</b>", body_style),
                Paragraph(f"{job['start_date']} – {job['end_date']}",
                          ParagraphStyle("DateR", parent=styles["Normal"],
                                         fontSize=9, alignment=TA_RIGHT))
            ],
            [
                Paragraph(f"<i>{job['company']}</i>", small_style),
                ""
            ]
        ]
        exp_table = Table(exp_data, colWidths=[page_w * 0.65, page_w * 0.35])
        exp_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
        ]))
        elements.append(exp_table)
        for bullet in job["bullets"]:
            elements.append(Paragraph(f"  • {bullet}", small_style))
        elements.append(Spacer(1, 4))

    # Education table
    edu_header = Table(
        [[Paragraph("EDUCATION", section_style)]],
        colWidths=[page_w]
    )
    edu_header.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#16213e")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]))
    elements.append(edu_header)
    edu_data = [[
        Paragraph(f"<b>{candidate['education']}</b> — {candidate['university']}", body_style),
        Paragraph(str(candidate["grad_year"]),
                  ParagraphStyle("YearR", parent=styles["Normal"],
                                 fontSize=9, alignment=TA_RIGHT))
    ]]
    edu_table = Table(edu_data, colWidths=[page_w * 0.75, page_w * 0.25])
    edu_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elements.append(edu_table)

    doc.build(elements)


# ---------------------------------------------------------------------------
# DOCX Generator
# ---------------------------------------------------------------------------

def _create_docx(candidate: dict, filepath: str):
    """Generate a DOCX resume."""
    doc = Document()

    # Adjust default style
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(10)

    # Name
    name_para = doc.add_heading(candidate["name"], level=1)
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in name_para.runs:
        run.font.color.rgb = RGBColor(26, 26, 46)

    # Contact line
    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_run = contact.add_run(
        f"{candidate['email']}  |  {candidate['phone']}  |  "
        f"{candidate['city']}, {candidate['state']}  |  {candidate['linkedin']}"
    )
    contact_run.font.size = Pt(9)
    contact_run.font.color.rgb = RGBColor(128, 128, 128)

    # Summary
    doc.add_heading("Professional Summary", level=2)
    doc.add_paragraph(candidate["summary"])

    # Skills
    doc.add_heading("Skills", level=2)
    skills_text = ", ".join(candidate["skills"])
    doc.add_paragraph(skills_text)

    # Experience
    doc.add_heading("Work Experience", level=2)
    for job in candidate["work_experience"]:
        job_para = doc.add_paragraph()
        title_run = job_para.add_run(f"{job['title']} — {job['company']}")
        title_run.bold = True
        job_para.add_run(f"  ({job['start_date']} – {job['end_date']})")

        for bullet in job["bullets"]:
            bullet_para = doc.add_paragraph(bullet, style="List Bullet")
            for run in bullet_para.runs:
                run.font.size = Pt(9)

    # Education
    doc.add_heading("Education", level=2)
    edu_para = doc.add_paragraph()
    edu_run = edu_para.add_run(f"{candidate['education']} — {candidate['university']}")
    edu_run.bold = True
    edu_para.add_run(f" ({candidate['grad_year']})")

    # Certifications
    if candidate["certifications"]:
        doc.add_heading("Certifications", level=2)
        for cert in candidate["certifications"]:
            doc.add_paragraph(cert, style="List Bullet")

    doc.save(filepath)


# ---------------------------------------------------------------------------
# Main Generator
# ---------------------------------------------------------------------------

def generate_resumes(output_dir: str, count: int = 200):
    """Generate synthetic resumes and save as PDF/DOCX."""
    os.makedirs(output_dir, exist_ok=True)

    pdf_generators = [
        ("single_col", _create_pdf_single_column),
        ("two_col", _create_pdf_two_column),
        ("table", _create_pdf_table_layout),
    ]

    candidates_metadata = []

    for i in range(1, count + 1):
        candidate = generate_candidate_data(i)

        # ~70% PDF, ~30% DOCX
        if random.random() < 0.7:
            # Pick a PDF layout variant
            layout_name, generator_fn = random.choice(pdf_generators)
            filename = f"resume_{i:04d}_{layout_name}.pdf"
            filepath = os.path.join(output_dir, filename)
            try:
                generator_fn(candidate, filepath)
            except Exception as e:
                print(f"  [WARN] PDF generation failed for candidate {i} "
                      f"(layout={layout_name}): {e}")
                # Fallback to single column
                filename = f"resume_{i:04d}_single_col.pdf"
                filepath = os.path.join(output_dir, filename)
                _create_pdf_single_column(candidate, filepath)
        else:
            filename = f"resume_{i:04d}.docx"
            filepath = os.path.join(output_dir, filename)
            _create_docx(candidate, filepath)

        candidate["filename"] = filename
        candidates_metadata.append(candidate)

        if i % 25 == 0:
            print(f"  Generated {i}/{count} resumes...")

    # Save metadata for downstream use (not the resumes themselves, but the
    # ground-truth data used to generate them — useful for evaluation)
    metadata_path = os.path.join(output_dir, "..", "candidates_metadata.json")
    with open(metadata_path, "w") as f:
        json.dump(candidates_metadata, f, indent=2, default=str)

    print(f"\n[OK] Generated {count} resumes in {output_dir}")
    print(f"  Metadata saved to {metadata_path}")
    return candidates_metadata


if __name__ == "__main__":
    output_dir = os.path.join(os.path.dirname(__file__), "..", "data", "raw_resumes")
    generate_resumes(output_dir, count=200)
