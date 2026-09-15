"""
Unit Tests for Parser Module
==============================
Tests:
  - PDF parsing returns valid text and tables
  - PyMuPDF fallback logic activates when pdfplumber extracts < 20 words
  - DOCX parsing extracts paragraphs and tables
"""

import os
import tempfile
import pytest
from unittest.mock import patch, MagicMock

from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document

from etl.parser import (
    parse_pdf, parse_docx, parse_file, MIN_WORD_COUNT,
    parse_pdf_pdfplumber, parse_pdf_pymupdf
)


@pytest.fixture
def sample_pdf(tmp_path):
    """Create a temporary PDF file with plenty of text."""
    pdf_path = tmp_path / "test_sample.pdf"
    c = canvas.Canvas(str(pdf_path), pagesize=letter)
    text = (
        "Alice Johnson is a Senior Software Engineer with 8 years of experience "
        "in Python, Django, PostgreSQL, Docker, and Kubernetes. "
        "She has designed and maintained microservices architectures serving millions of users."
    )
    c.drawString(100, 700, text)
    c.save()
    return str(pdf_path)


@pytest.fixture
def empty_pdf(tmp_path):
    """Create a PDF with fewer than MIN_WORD_COUNT words."""
    pdf_path = tmp_path / "test_empty.pdf"
    c = canvas.Canvas(str(pdf_path), pagesize=letter)
    c.drawString(100, 700, "Short resume.")
    c.save()
    return str(pdf_path)


@pytest.fixture
def sample_docx(tmp_path):
    """Create a temporary DOCX file with text and tables."""
    docx_path = tmp_path / "test_sample.docx"
    doc = Document()
    doc.add_heading("Bob Smith", level=1)
    doc.add_paragraph("Data Analyst with 4 years of experience in SQL and Python.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "Proficiency"
    table.cell(1, 0).text = "SQL"
    table.cell(1, 1).text = "Advanced"

    doc.save(str(docx_path))
    return str(docx_path)


def test_parse_pdf_success(sample_pdf):
    """Test standard PDF parsing returns non-empty text."""
    result = parse_pdf(sample_pdf)
    assert "raw_text" in result
    assert "tables" in result
    assert len(result["raw_text"].split()) >= 10
    assert "Alice Johnson" in result["raw_text"]


def test_parse_pdf_fallback_trigger(empty_pdf):
    """Test that PyMuPDF fallback is triggered when pdfplumber yields < MIN_WORD_COUNT words."""
    with patch("etl.parser.parse_pdf_pdfplumber") as mock_pdfplumber, \
         patch("etl.parser.parse_pdf_pymupdf") as mock_pymupdf:

        # pdfplumber returns only 5 words (< 20 threshold)
        mock_pdfplumber.return_value = {"raw_text": "Short note only five words", "tables": []}
        mock_pymupdf.return_value = {"raw_text": "Fallback text from fitz", "tables": []}

        result = parse_pdf(empty_pdf)

        mock_pdfplumber.assert_called_once_with(empty_pdf)
        mock_pymupdf.assert_called_once_with(empty_pdf)
        assert result["raw_text"] == "Fallback text from fitz"


def test_parse_docx_success(sample_docx):
    """Test that DOCX parsing extracts text and structured table data."""
    result = parse_docx(sample_docx)
    assert "Bob Smith" in result["raw_text"]
    assert "Data Analyst" in result["raw_text"]
    assert len(result["tables"]) == 1
    assert result["tables"][0][1][0] == "SQL"
    assert result["tables"][0][1][1] == "Advanced"


def test_parse_file_dispatcher(sample_pdf, sample_docx):
    """Test parse_file correctly handles both .pdf and .docx extensions."""
    pdf_res = parse_file(sample_pdf)
    assert pdf_res["source_file"] == os.path.basename(sample_pdf)

    docx_res = parse_file(sample_docx)
    assert docx_res["source_file"] == os.path.basename(sample_docx)


def test_parse_file_unsupported_format(tmp_path):
    """Test unsupported extensions raise ValueError."""
    txt_path = tmp_path / "test.txt"
    txt_path.write_text("plain text")

    with pytest.raises(ValueError, match="Unsupported file format"):
        parse_file(str(txt_path))
